# -*- coding: utf-8 -*-
"""Chèn mục 'Mã nguồn chức năng' vào sau mục 'Màn hình ...' của 13 chức năng
trong tài liệu Word: 05_Bản mô tả Ứng dụng di động.

Mỗi mục mã nguồn gồm 3 khối: 25 dòng đầu / 25 dòng giữa / 25 dòng cuối
trích từ phân đoạn đại diện của file src/App.jsx.
"""
import copy
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

ROOT = Path(r"D:/Etc/ITS_Mobile_VEC")
SRC_DOC = ROOT / "docs/source/05_Bản_mô_tả_Ứng_dụng_di_động_cho_nhân_viên_vận_hành_tuyến_cao_tốc.docx"
import os
# Nếu file đang bị Word lock, ghi sang file tạm có hậu tố _v2
def _resolve_out():
    out = SRC_DOC
    try:
        with open(out, "ab"):
            pass
        return out
    except PermissionError:
        return out.with_name(out.stem + "_v2" + out.suffix)
OUT_DOC = _resolve_out()
APP_JSX = ROOT / "src/App.jsx"

# Phạm vi mã nguồn đại diện cho từng chức năng (1-indexed, inclusive).
# Mỗi phạm vi phải đủ rộng để chứa 25 dòng đầu + 25 dòng giữa + 25 dòng cuối
# (tối thiểu 75 dòng; có thể overlap nhẹ với chức năng lân cận khi component nằm sát nhau).
FUNCTIONS = [
    # (heading H3 text, file relative, line_start, line_end)
    ("Chức năng Đăng nhập",                "src/App.jsx", 359, 451),  # LoginScreen
    ("Chức năng Đăng xuất",                "src/App.jsx", 1050, 1140), # Profile tab + nút Đăng xuất
    ("Thực hiện cuộc gọi nội bộ PBX",      "src/App.jsx", 833, 921),  # CallsView
    ("Gọi khẩn cấp SOS qua GSM",           "src/App.jsx", 1100, 1180), # Nút SOS + modal xác nhận
    ("Hiển thị thông tin cá nhân",         "src/App.jsx", 452, 530),  # App init + session restore
    ("Chức năng Hiển thị danh sách công việc", "src/App.jsx", 681, 757),  # TaskListView (phần Nhiệm vụ)
    ("Chức năng Xem chi tiết công việc",   "src/App.jsx", 115, 200),  # TaskDetailViewPanel (header + data)
    ("Cập nhật trạng thái xử lý sự cố",    "src/App.jsx", 220, 300),  # applyStatus + status timeline
    ("Đính kèm ảnh/ video hiện trường",    "src/App.jsx", 280, 356),  # Attach UI + handlers
    ("Hiển thị sự cố/ sự kiện giao thông", "src/App.jsx", 745, 832),  # Events trên tuyến
    ("Nhận Push Notification",             "src/App.jsx", 552, 630),  # notifications data + unreadCount
    ("Chức năng xem danh sách thông báo",  "src/App.jsx", 922, 1006), # NotificationsView
    ("Đánh dấu đã đọc/ chưa đọc",          "src/App.jsx", 540, 625),  # notifications array với cờ unread
]

# ───────────────────────── helpers ─────────────────────────

def read_lines(path: Path):
    """Đọc nguyên file, trả về list dòng đã strip newline cuối."""
    return path.read_text(encoding="utf-8").splitlines()


def slice_three_chunks(lines, start, end):
    """Trả về (first25, middle25, last25) — mỗi cái là list[(line_no, text)].

    start/end: 1-indexed inclusive.
    """
    # convert sang 0-indexed half-open
    s0 = start - 1
    e0 = end  # exclusive
    block = lines[s0:e0]
    block_lines = [(s0 + 1 + i, t) for i, t in enumerate(block)]
    n = len(block_lines)

    first = block_lines[:25]
    last = block_lines[max(0, n - 25):]
    mid_start = max(0, (n - 25) // 2)
    middle = block_lines[mid_start:mid_start + 25]
    return first, middle, last


def _code_run(text):
    """Tạo 1 <w:r> chứa text mã nguồn (Consolas 8pt, xám đậm, giữ space)."""
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Consolas")
    rFonts.set(qn("w:hAnsi"), "Consolas")
    rFonts.set(qn("w:cs"), "Consolas")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")  # 8pt
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "16")
    rPr.append(szCs)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F2933")
    rPr.append(color)
    r.append(rPr)

    t = OxmlElement("w:t")
    safe = text.replace("\t", "  ")
    t.text = safe if safe else " "
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def make_code_paragraphs(doc, chunk):
    """Trả về list <w:p>: mỗi dòng code = 1 paragraph riêng để tránh Word
    'justify' bóp giãn các dòng (nguyên nhân khoảng trắng khổng lồ giữa từ).

    - Paragraph đầu tiên: border top + left + right
    - Paragraph cuối:    border bottom + left + right
    - Paragraph giữa:    border left + right
    Tất cả: shading nền F4F4F4, jc=left, font Consolas 8pt, line-spacing chặt.
    """
    results = []
    n = len(chunk)
    for i, (_lineno, text) in enumerate(chunk):
        p = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")

        # alignment: TRÁI (ngăn Word justify bóp giãn space)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "left")
        pPr.append(jc)

        # border: chỉ top ở dòng đầu, bottom ở dòng cuối, left+right cho mọi dòng
        pBdr = OxmlElement("w:pBdr")
        sides = ["left", "right"]
        if i == 0:
            sides.append("top")
        if i == n - 1:
            sides.append("bottom")
        for side in sides:
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "4")
            b.set(qn("w:color"), "BFBFBF")
            pBdr.append(b)
        pPr.append(pBdr)

        # shading
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F4F4F4")
        pPr.append(shd)

        # spacing: giữa các dòng = 0, only before ở dòng đầu / after ở dòng cuối
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "60" if i == 0 else "0")
        spacing.set(qn("w:after"), "120" if i == n - 1 else "0")
        spacing.set(qn("w:line"), "220")
        spacing.set(qn("w:lineRule"), "auto")
        pPr.append(spacing)

        # giữ chung block
        pPr.append(OxmlElement("w:keepLines"))
        if i < n - 1:
            pPr.append(OxmlElement("w:keepNext"))

        # không thụt lề
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "0")
        ind.set(qn("w:right"), "0")
        ind.set(qn("w:firstLine"), "0")
        pPr.append(ind)

        p.append(pPr)
        p.append(_code_run(text))
        results.append(p)
    return results


def make_styled_paragraph(text, style_name=None, bold=False, italic=False, size_pt=None, color_hex=None):
    """Tạo <w:p> với 1 run text và style name (Heading 4, v.v.)."""
    p = OxmlElement("w:p")
    if style_name:
        pPr = OxmlElement("w:pPr")
        pStyle = OxmlElement("w:pStyle")
        # python-docx lưu style id; với style mặc định Heading 4 id = "Heading4"
        style_id = style_name.replace(" ", "")
        pStyle.set(qn("w:val"), style_id)
        pPr.append(pStyle)
        p.append(pPr)

    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if bold:
        rPr.append(OxmlElement("w:b"))
    if italic:
        rPr.append(OxmlElement("w:i"))
    if size_pt:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size_pt * 2)))
        rPr.append(sz)
    if color_hex:
        col = OxmlElement("w:color")
        col.set(qn("w:val"), color_hex)
        rPr.append(col)
    if rPr.findall("*"):
        r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    p.append(r)
    return p


def find_function_h3_indices(doc):
    """Trả về dict {heading_text: index} cho 13 chức năng."""
    body_paras = list(doc.paragraphs)
    result = {}
    wanted = {f[0] for f in FUNCTIONS}
    for i, p in enumerate(body_paras):
        if p.style and p.style.name == "Heading 3":
            txt = p.text.strip()
            if txt in wanted and txt not in result:
                result[txt] = i
    return result


def insert_block_before_paragraph(target_para, elements):
    """Chèn các <w:p> trước paragraph target_para (python-docx Paragraph)."""
    target_xml = target_para._p
    for el in elements:
        target_xml.addprevious(el)


def append_elements_to_body(doc, elements):
    """Chèn các <w:p> vào cuối body trước sectPr."""
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    if sectPr is not None:
        for el in elements:
            sectPr.addprevious(el)
    else:
        for el in elements:
            body.append(el)


# ───────────────────────── main ─────────────────────────

def build_block_for_function(doc, file_rel, start, end):
    """Trả về list các <w:p> tạo nên block 'Mã nguồn chức năng' cho 1 function."""
    lines = read_lines(ROOT / file_rel)
    first, middle, last = slice_three_chunks(lines, start, end)

    parts = []
    # H4 heading
    parts.append(make_styled_paragraph("Mã nguồn chức năng", style_name="Heading 4"))
    # File đại diện
    parts.append(make_styled_paragraph(
        f"File đại diện: {file_rel} (phân đoạn dòng {start}–{end}, tổng {end - start + 1} dòng).",
        italic=True, size_pt=10, color_hex="555555",
    ))
    # First 25
    first_lo = first[0][0] if first else start
    first_hi = first[-1][0] if first else start
    parts.append(make_styled_paragraph(
        f"• 25 dòng đầu (dòng {first_lo}–{first_hi}):",
        bold=True, size_pt=10, color_hex="0B5394",
    ))
    parts.extend(make_code_paragraphs(doc, first))
    # Middle 25
    mid_lo = middle[0][0] if middle else start
    mid_hi = middle[-1][0] if middle else start
    parts.append(make_styled_paragraph(
        f"• 25 dòng giữa (dòng {mid_lo}–{mid_hi}):",
        bold=True, size_pt=10, color_hex="0B5394",
    ))
    parts.extend(make_code_paragraphs(doc, middle))
    # Last 25
    last_lo = last[0][0] if last else start
    last_hi = last[-1][0] if last else start
    parts.append(make_styled_paragraph(
        f"• 25 dòng cuối (dòng {last_lo}–{last_hi}):",
        bold=True, size_pt=10, color_hex="0B5394",
    ))
    parts.extend(make_code_paragraphs(doc, last))
    return parts


def remove_existing_injections(doc):
    """Xoá các block 'Mã nguồn chức năng' đã chèn ở lần trước.
    Quy tắc: tìm mọi paragraph H4 có text == 'Mã nguồn chức năng' rồi xoá nó +
    các paragraph kế tiếp cho tới (nhưng không bao gồm) paragraph có style
    Heading 1/2/3, hoặc cho tới hết body.
    """
    body = doc.element.body
    # Lấy danh sách paragraph elements (chỉ paragraph trực tiếp dưới body)
    p_elems = body.findall(qn("w:p"))
    to_delete = []
    i = 0
    removed_blocks = 0
    while i < len(p_elems):
        p = p_elems[i]
        # detect style + text
        pStyle = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        style_id = pStyle.get(qn("w:val")) if pStyle is not None else ""
        texts = [t.text or "" for t in p.findall(".//" + qn("w:t"))]
        full_text = "".join(texts).strip()

        if style_id == "Heading4" and full_text == "Mã nguồn chức năng":
            # xoá block này
            to_delete.append(p)
            j = i + 1
            while j < len(p_elems):
                q = p_elems[j]
                qStyle = q.find(qn("w:pPr") + "/" + qn("w:pStyle"))
                qStyleId = qStyle.get(qn("w:val")) if qStyle is not None else ""
                if qStyleId in ("Heading1", "Heading2", "Heading3"):
                    break
                to_delete.append(q)
                j += 1
            removed_blocks += 1
            i = j
        else:
            i += 1

    for el in to_delete:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
    print(f"Removed {removed_blocks} existing 'Mã nguồn chức năng' block(s) ({len(to_delete)} paragraphs).")


def main():
    doc = Document(SRC_DOC)
    remove_existing_injections(doc)
    h3_idx = find_function_h3_indices(doc)
    print(f"Found {len(h3_idx)} H3 headings: {list(h3_idx.keys())}")

    body_paras = list(doc.paragraphs)

    # Tính position chèn cho từng function: TRƯỚC paragraph "ngắt section" kế tiếp.
    # Paragraph ngắt = H2 hoặc H3 (lấy cái xuất hiện trước, tính từ H3 hiện tại).
    # Như vậy, nếu giữa 2 H3 có 1 H2 chen vào (đầu nhóm chức năng mới), khối mã nguồn
    # vẫn nằm gọn TRONG H3 hiện tại và TRƯỚC H2 đó.
    # Function cuối (index 12) chèn vào CUỐI body.
    insert_jobs = []  # list[(insert_before_para_or_None, function_data)]

    for fi, (h3_text, file_rel, start, end) in enumerate(FUNCTIONS):
        if h3_text not in h3_idx:
            raise SystemExit(f"❌ Không tìm thấy heading H3: {h3_text!r}")
        current_idx = h3_idx[h3_text]
        if fi + 1 < len(FUNCTIONS):
            # tìm paragraph đầu tiên có style H2 hoặc H3 sau current_idx
            target_para = None
            for j in range(current_idx + 1, len(body_paras)):
                style = body_paras[j].style.name if body_paras[j].style else ""
                if style in ("Heading 2", "Heading 3"):
                    target_para = body_paras[j]
                    break
            if target_para is None:
                raise SystemExit(f"❌ Không tìm thấy H2/H3 kế tiếp sau {h3_text!r}")
        else:
            target_para = None  # append to body end
        insert_jobs.append((target_para, (file_rel, start, end), h3_text))

    # Thực hiện chèn — vì chèn trước paragraph cụ thể nên thứ tự không ảnh hưởng index gốc
    for target_para, (file_rel, start, end), h3_text in insert_jobs:
        block = build_block_for_function(doc, file_rel, start, end)
        if target_para is not None:
            insert_block_before_paragraph(target_para, block)
            print(f"  ✓ Inserted block for {h3_text!r} before next H3")
        else:
            append_elements_to_body(doc, block)
            print(f"  ✓ Appended block for {h3_text!r} at end of body")

    doc.save(OUT_DOC)
    print(f"\n✅ Saved: {OUT_DOC}")


if __name__ == "__main__":
    main()
